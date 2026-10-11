# 연습 업무카드에 붙일 빈 양식(.docx) 2개를 만든다 — 집행기준(예규 제372호) 서식의 항목을 따른 빈 칸 양식
# 실행: (tools 폴더에서) python make_practice_forms.py  → practice/ 폴더
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'practice')
os.makedirs(OUT, exist_ok=True)


def new_doc():
    d = Document()
    st = d.styles['Normal']
    st.font.name = '맑은 고딕'
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    st.font.size = Pt(11)
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(2.5)
    return d


def title(d, text):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(18)
    d.add_paragraph()


def table(d, rows, widths):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.width = Cm(widths[j])
            c.text = v
            if j == 0 or i == 0 and len(rows[0]) > 2:
                for r in c.paragraphs[0].runs:
                    r.bold = True
    return t


# 1) 각서 — 집행기준 제5장 [별지 제1호 서식]의 항목
d = new_doc()
title(d, '각    서')
table(d, [['업 체 명', ''], ['대 표 자', ''], ['소 재 지', ''], ['업종(등록)', '']], [4, 12])
d.add_paragraph()
d.add_paragraph('상기 본인(법인)은 귀 기관과 수의계약을 체결함에 있어서 붙임 배제사유 중 어느 사유에도 '
                '해당되지 않으며, 차후에 이러한 사실이 발견된 경우 계약의 해제·해지 및 부정당업자 제재 처분을 '
                '받아도 이의를 제기하지 않겠습니다.')
d.add_paragraph()
d.add_paragraph('붙임 : 수의계약 배제사유 1부.')
d.add_paragraph()
p = d.add_paragraph('20    .    .    .')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = d.add_paragraph('업체명 :                    \n대표자 :                    (인)')
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
d.add_paragraph()
d.add_paragraph('재무관 또는 분임재무관 귀하')
d.add_paragraph()
n = d.add_paragraph('※ 연습용 빈 양식입니다. 실제 업무에는 최신 「지방자치단체 입찰 및 계약 집행기준」 제5장 [별지 제1호 서식]을 확인하세요.')
n.runs[0].font.size = Pt(9)
d.save(os.path.join(OUT, '각서(수의계약 배제사유).docx'))

# 2) 수의계약 공개내역서 — 집행기준 제5장 제8절의 공개내역서 항목
d = new_doc()
title(d, '수의계약 공개내역서')
table(d, [
    ['사 업 명', '', '', ''],
    ['계약 개요', '', '', ''],
    ['계약일자', '', '계약기간', ''],
    ['예정가격(추정금액)(A)', '', '계약금액(B)', ''],
    ['계약률(%) (B/A)', '', '', ''],
    ['계약상대자', '업체명', '', ''],
    ['', '대표자', '', ''],
    ['', '주소', '', ''],
    ['수의계약 사유', '※ 관련 법령 근거 및 구체적인 사유를 반드시 적음', '', ''],
    ['사업 장소', '', '', ''],
    ['기    타', '', '', ''],
], [4.2, 4, 3.6, 4.2])
d.add_paragraph()
n = d.add_paragraph('※ 연습용 빈 양식입니다. 실제 업무에는 최신 「지방자치단체 입찰 및 계약 집행기준」 제5장 제8절을 확인하세요.')
n.runs[0].font.size = Pt(9)
d.save(os.path.join(OUT, '수의계약 공개내역서.docx'))
print('practice/ 에 양식 2개')
